from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# Define monospace font style for diagrams
style_mono = doc.styles.add_style('Monospace', 1)
style_mono.font.name = 'Courier New'
style_mono.font.size = Pt(8)

# Title
title = doc.add_heading('Nebula Kirana Agent', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

subtitle = doc.add_paragraph("Run an entire Indian kirana store from a chat window — with an agent, not a menu.\n")
subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
subtitle.runs[0].italic = True

links = doc.add_paragraph()
links.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
r1 = links.add_run("Live Bot: https://t.me/NebulaKirana_Saravanan_bot\n")
r1.bold = True
r2 = links.add_run("GitHub Repository: https://github.com/Saravanan460/nebula-kirana-agent.git")
r2.bold = True

doc.add_page_break()

# Overview
doc.add_heading('🌟 Overview', level=1)
doc.add_paragraph("The Nebula Kirana Agent completely replaces complex POS software. The store owner runs their entire shop simply by chatting with this agent on Telegram — no menus, no forms, no admin panel. The chat is the product.")

features = [
    "📦 Inventory Management — add products, receive stock, search, low-stock alerts, list entire inventory",
    "🧾 GST-Compliant Billing — multi-item draft bills, mid-build edits, finalization with tax breakup",
    "📒 Khata (Credit Ledger) — create customer credit, charge bills, record payments, check balances",
    "📊 Analytics & Documents — PDF invoices (reportlab) & Weekly PPTX sales decks (python-pptx) with real charts",
    "🧠 Persistent Memory — owner preferences survive /new chat resets",
    "🏢 Multi-Tenant — every Telegram user gets a completely isolated private store sandbox"
]
for f in features:
    doc.add_paragraph(f, style='List Bullet')

# Architecture Diagram
doc.add_heading('🏗️ Architecture', level=1)
arch_diagram = """
┌──────────────────────────────────────────────────────────────────┐
│                        TELEGRAM USER                             │
│            "make a bill: 2kg sugar, 4 maggi, UPI"                │
└──────────────────────┬───────────────────────────────────────────┘
                       │
                       ▼
┌──────────────────────────────────────────────────────────────────┐
│                   TELEGRAM BOT LAYER                             │
│  handlers.py — routes /start, /new, and natural language msgs    │
│  Maintains per-chat conversation history (last 20 messages)      │
│  Converts **bold** → <b>HTML</b> for clean Telegram rendering   │
└──────────────────────┬───────────────────────────────────────────┘
                       │
                       ▼
┌──────────────────────────────────────────────────────────────────┐
│                  PYDANTIC-AI AGENT (harness.py)                  │
│                                                                  │
│  ┌─────────────┐    The LLM receives the user's natural          │
│  │  Groq API   │    language, reasons over it, and decides       │
│  │ (gpt-oss)   │    which tools to call. It chains multiple      │
│  │     or      │    tool calls in a single turn as needed.       │
│  │ Gemini 2.0  │    Business rules live in the TOOLS,            │
│  │  (fallback) │    not the prompt.                              │
│  └─────────────┘                                                 │
│                                                                  │
│  System Prompt dynamically injects owner preferences             │
│  from SQLite on every single message.                            │
└──────────────────────┬───────────────────────────────────────────┘
                       │ Calls tools based on reasoning
                       ▼
┌──────────────────────────────────────────────────────────────────┐
│                     TOOL / SKILL LAYER                           │
│                                                                  │
│  ┌────────────┐ ┌────────────┐ ┌────────────┐ ┌──────────────┐  │
│  │ inventory  │ │  billing   │ │   khata    │ │ preferences  │  │
│  │            │ │            │ │            │ │              │  │
│  │ add_product│ │ start_bill │ │create_khata│ │set_preference│  │
│  │recv_stock  │ │ add_item   │ │charge_khata│ │get_preference│  │
│  │check_stock │ │ edit_item  │ │ pay_khata  │ │              │  │
│  │search      │ │ view_bill  │ │check_khata │ │              │  │
│  │list_all    │ │ finalize   │ │            │ │              │  │
│  │low_stock   │ │ cancel     │ │            │ │              │  │
│  └────────────┘ └────────────┘ └────────────┘ └──────────────┘  │
│                                                                  │
│  ┌────────────┐ ┌──────────────────────────────────────────────┐ │
│  │ analytics  │ │           documents                          │ │
│  │            │ │                                              │ │
│  │daily_sum   │ │ generate_invoice_pdf  (reportlab)            │ │
│  │close_day   │ │ generate_analysis_pptx (python-pptx + charts)│ │
│  └────────────┘ └──────────────────────────────────────────────┘ │
└──────────────────────┬───────────────────────────────────────────┘
                       │ All tools enforce business rules
                       ▼
┌──────────────────────────────────────────────────────────────────┐
│                   SQLite DATABASE (WAL mode)                     │
│                                                                  │
│  ┌──────────┐ ┌───────┐ ┌───────────┐ ┌───────┐ ┌───────────┐  │
│  │ products │ │ bills │ │bill_items │ │ khata │ │preferences│  │
│  │(per user)│ │       │ │           │ │       │ │           │  │
│  └──────────┘ └───────┘ └───────────┘ └───────┘ └───────────┘  │
│                                                                  │
│  Every table is scoped by chat_id → full multi-tenant isolation  │
│  Data survives restarts. /new clears context, not the database.  │
└──────────────────────────────────────────────────────────────────┘
"""
p_arch = doc.add_paragraph(arch_diagram.strip(), style='Monospace')

doc.add_heading('Control Loop', level=2)
doc.add_paragraph("The agent follows a strict Observe → Reason → Act → Feed Back loop:")
loop_diagram = """
User: "make a bill: 2kg sugar and 4 maggi, UPI"
  │
  ├─→ LLM reasons: "I need SKUs. Let me search first."
  │     ├─→ TOOL CALL: search_product("sugar")     → finds SUGAR-LOOSE
  │     ├─→ TOOL CALL: search_product("maggi")     → finds MAGGI-70G
  │     ├─→ TOOL CALL: start_bill()                → draft created
  │     ├─→ TOOL CALL: add_item("SUGAR-LOOSE", 2)  → ✅ added
  │     ├─→ TOOL CALL: add_item("MAGGI-70G", 4)    → ✅ added
  │     └─→ TOOL CALL: finalize_bill("upi")        → ✅ stock decremented
  │
  └─→ Agent replies: "Bill finalized! Total ₹146 via UPI."
"""
p_loop = doc.add_paragraph(loop_diagram.strip(), style='Monospace')

doc.add_page_break()

# Tech Stack Table
doc.add_heading('🛠️ The 100% Free Tech Stack', level=1)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Component'
hdr_cells[1].text = 'Choice'
hdr_cells[2].text = 'Why'

tech_stack = [
    ('Agent Harness', 'pydantic-ai', 'Clean, type-safe, agent-first — no LangGraph-style node boilerplate'),
    ('LLM Engine', 'Groq API (gpt-oss-120b)', 'Lightning-fast tool-calling; fallback to Gemini 2.0 Flash'),
    ('Database', 'SQLite (WAL mode)', 'Built into Python, free, durable, handles concurrency'),
    ('Interface', 'Telegram Bot API', 'Zero cost, file sharing (PDF/PPTX), accessible to shopkeepers'),
    ('PDF Generation', 'reportlab', 'Offline GST-compliant invoice generation'),
    ('PPTX Generation', 'python-pptx', 'Native charts and slides for business analysis decks'),
    ('Hosting', 'PythonAnywhere', 'Free 24/7 cloud hosting for the live bot')
]

for item in tech_stack:
    row_cells = table.add_row().cells
    row_cells[0].text = item[0]
    row_cells[1].text = item[1]
    row_cells[2].text = item[2]

# Edge Cases & Guardrails
doc.add_heading('🧠 Engineering Solutions to the "Hard Parts"', level=1)

edge_cases = [
    ("🛡️ 1. Grounding & Oversell Guard", "Prices, GST slabs and stock come from the DB via tools. Never invent a product or a price. Stock validation happens in the tool/database layer, blocking oversell explicitly. finalize_bill checks that no item's MRP is below its cost_price."),
    ("🔄 2. Concurrency & Race Conditions", "WAL mode enabled on SQLite startup. BEGIN IMMEDIATE locks the database during critical stock modifications. Atomic decrement is the absolute final guard against negative stock."),
    ("🧾 3. Idempotency (No Double-Billing)", "Telegram redelivers updates. The Telegram message_id is used as a unique txn_id. finalize_bill enforces a UNIQUE(txn_id) constraint via SAVEPOINT, preventing double billing on retries."),
    ("💬 4. Stateful Multi-Turn Bills", "A bill builds over several messages, supports edits, and only decrements stock on finalize. The bills table tracks active drafts."),
    ("📊 5. Real Artifact Generation", "PDF Invoices via reportlab format standard Tax Invoices with HSN codes and SGST/CGST split. Weekly Decks via python-pptx query SQLite for top-selling items and payment mode breakdowns."),
    ("🧠 6. Durable Memory Across Sessions", "Preferences are saved to a durable preferences table. /new wipes conversation context but preferences are dynamically injected into the system prompt on every message."),
    ("🏢 7. Multi-Tenant Sandbox Architecture", "The entire database is strictly isolated by chat_id. Auto-Seeding initializes mandatory SKUs on /start for every user."),
    ("✅ 8. GST Correctness", "Each product carries its own gst_rate and hsn_code. MRP is treated as GST-inclusive. The tool back-calculates base amount, splits into CGST + SGST, and rounds correctly.")
]

for title, desc in edge_cases:
    doc.add_heading(title, level=2)
    doc.add_paragraph(desc)

doc.add_page_break()

# Deployment
doc.add_heading('🚀 Deployment (PythonAnywhere)', level=1)
doc.add_paragraph("1. Clone the GitHub repository:\n    git clone https://github.com/Saravanan460/nebula-kirana-agent.git")
doc.add_paragraph("2. Setup virtual environment:\n    python -m venv venv\n    source venv/bin/activate\n    pip install -r requirements.txt")
doc.add_paragraph("3. Configure Environment (.env):\n    TELEGRAM_BOT_TOKEN=your_telegram_bot_token\n    GROQ_API_KEY=your_groq_api_key\n    GEMINI_API_KEY=your_gemini_api_key")
doc.add_paragraph("4. Run the Agent:\n    python main.py")

doc.save('Nebula_Kirana_Agent_Project_Document.docx')
print("Document successfully regenerated with diagrams and tables.")
